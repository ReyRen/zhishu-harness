"""Create and reopen Office files with the image's bundled authoring libraries."""

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

root = Path(sys.argv[1])
text = "Office runtime 中文验证"
assert int(np.arange(4).sum()) == 6
assert int(pd.DataFrame({"n": [1, 2]}).n.sum()) == 3

document = Document()
document.add_heading(text, 0)
document.add_paragraph("Word authoring works offline.")
document.save(root / "sample.docx")
assert text in "\n".join(p.text for p in Document(root / "sample.docx").paragraphs)

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = text
slide.placeholders[1].text = "PowerPoint authoring works offline."
presentation.save(root / "sample.pptx")
assert Presentation(root / "sample.pptx").slides[0].shapes.title.text == text

workbook = Workbook()
sheet = workbook.active
sheet.title = "Report"
sheet["A1"] = text
sheet["A2"], sheet["B2"] = 2, 3
sheet["C2"] = "=SUM(A2:B2)"
workbook.save(root / "sample.xlsx")
assert load_workbook(root / "sample.xlsx").active["C2"].value == "=SUM(A2:B2)"
print("Python libraries and DOCX/PPTX/XLSX read/write passed")
